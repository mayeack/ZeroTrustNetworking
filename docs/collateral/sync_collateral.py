#!/usr/bin/env python3
"""make collateral: apply docs/collateral/collateral.yaml to the deck and the talk track in the OneDrive folder.
--dry-run prints the changes; --live fills the `live:` section from the stack's last run first. A one-time backup of
each file is kept in docs/collateral/backup/."""
import copy
import hashlib
import os
import shutil
import sys

import yaml

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
BANNED = ("sample", "mockup", "mock ", "fake", "synthetic", "illustrative")


def backup(path):
    dst = os.path.join(HERE, "backup", os.path.basename(path))
    if not os.path.exists(dst):
        shutil.copy2(path, dst)
        print("backup -> %s" % dst)


# ----------------------------------------------------------------------------- deck
def sync_deck(cfg, dry):
    from pptx import Presentation
    path = cfg["files"]["deck"]
    prs = Presentation(path)
    slides = list(prs.slides)
    changed = 0
    for rep in cfg["deck"].get("text_replacements", []):
        slide = slides[rep["slide"] - 1]
        match = rep["match"].split(" | ")
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            paras = sh.text_frame.paragraphs
            texts = [p.text for p in paras]
            if texts[:len(match)] == match or sh.text_frame.text.replace("\n", " | ") == rep["match"]:
                new = rep["replace"].split(" | ")
                for p, t in zip(paras, new):
                    if p.runs:
                        p.runs[0].text = t
                        for r in p.runs[1:]:
                            r.text = ""
                    else:
                        p.text = t
                changed += 1
                print("slide %d: text '%s' -> '%s'" % (rep["slide"], rep["match"][:50], rep["replace"][:50]))
    for num, text in (cfg["deck"].get("notes") or {}).items():
        slide = slides[int(num) - 1]
        tf = slide.notes_slide.notes_text_frame
        if tf.text != text:
            tf.text = text
            changed += 1
            print("slide %s: notes updated (%d chars)" % (num, len(text)))
    for rep in cfg["deck"].get("picture_replacements", []):
        shot = os.path.join(HERE, cfg["screenshots"].get(rep["screenshot"], ""))
        if not os.path.exists(shot):
            print("slide %d: screenshot %s missing, picture kept" % (rep["slide"], shot))
            continue
        slide = slides[rep["slide"] - 1]
        want = hashlib.sha1(open(shot, "rb").read()).hexdigest()
        for sh in slide.shapes:
            if sh.shape_type == 13 and sh.name.startswith(rep["shape_prefix"]):
                if sh.image.sha1 == want:  # already this screenshot: nothing to do (keeps --dry-run at 0)
                    break
                left, top, width, height = sh.left, sh.top, sh.width, sh.height
                sp = sh._element
                sp.getparent().remove(sp)
                pic = slide.shapes.add_picture(shot, left, top, width=width)
                if pic.height > height:  # keep the box: crop the bottom
                    pic.crop_bottom = 1 - height / pic.height
                    pic.height = height
                pic.name = sh.name
                changed += 1
                print("slide %d: picture %s replaced with %s" % (rep["slide"], sh.name, os.path.basename(shot)))
                break
    check_words(prs)
    if changed and not dry:
        backup(path)
        prs.save(path)
        print("deck saved (%d changes)" % changed)
    return changed


def check_words(prs):
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.has_text_frame:
                low = sh.text_frame.text.lower()
                for w in BANNED:
                    if w in low:
                        print("WARNING slide %d shape %s contains '%s'" % (i, sh.name, w.strip()))


# ----------------------------------------------------------------------------- talk track
def sync_talk_track(cfg, dry):
    import docx
    path = cfg["files"]["talk_track"]
    doc = docx.Document(path)
    tt = cfg["talk_track"]
    changed = 0
    # section 3 table: columns "Today" and "Before the lab phase"; missing items are added as styled rows
    for table in doc.tables:
        header = [c.text.strip() for c in table.rows[0].cells]
        if header[:2] == ["Item", "Today"]:
            present = {r.cells[0].text.strip() for r in table.rows[1:]}
            for item in tt["section3_rows"]:
                if item not in present:
                    add_styled_row(table, [item, tt["section3_rows"][item], tt.get("section3_before", {}).get(item, "")])
                    changed += 1
                    print("talk track section 3: row %s added" % item)
            for row in table.rows[1:]:
                item = row.cells[0].text.strip()
                for col, key in ((1, "section3_rows"), (2, "section3_before")):
                    want = tt.get(key, {}).get(item)
                    if want is not None and row.cells[col].text.strip() != want:
                        set_cell(row.cells[col], want)
                        changed += 1
                        print("talk track section 3: row %s column %d updated" % (item, col + 1))
        if header[:2] == ["Claim", "Status and source"]:
            existing = {r.cells[0].text.strip() for r in table.rows[1:]}
            for claim, status in tt.get("section7_rows", []):
                if claim not in existing:
                    add_styled_row(table, [claim, status])
                    changed += 1
                    print("talk track section 7: row added: %s" % claim[:60])
    # beats 5-7: replace heading text and body until the next Heading2
    paras = doc.paragraphs
    heads = [(i, p) for i, p in enumerate(paras) if (p.style is not None and p.style.name.startswith("Heading 2")) and p.text.startswith("Beat ")]
    templates = find_templates(paras)
    for beat in tt["beats"]:
        n = beat["title"].split(":")[0]  # "Beat 5"
        for idx, (i, p) in enumerate(heads):
            if p.text.startswith(n + ":"):
                end = heads[idx + 1][0] if idx + 1 < len(heads) else next(j for j, q in enumerate(paras) if j > i and q.style is not None and q.style.name.startswith("Heading 1"))
                current = [q.text for q in paras[i + 1:end]]
                wanted = beat_paragraphs(beat)
                if p.text != beat["title"] or current != [w[1] for w in wanted]:
                    set_paragraph(p, beat["title"])
                    for q in paras[i + 1:end]:
                        q._element.getparent().remove(q._element)
                    anchor = p._element
                    for kind, text in wanted:
                        newp = copy.deepcopy(templates[kind]._element)
                        anchor.addnext(newp)
                        anchor = newp
                        if kind in ("who", "say"):
                            set_lead_text(newp, text)
                        else:
                            set_element_text(newp, text)
                    changed += 1
                    print("talk track: %s rewritten" % beat["title"])
                    paras = doc.paragraphs
                    heads = [(i2, p2) for i2, p2 in enumerate(paras) if (p2.style is not None and p2.style.name.startswith("Heading 2")) and p2.text.startswith("Beat ")]
                break
    if changed and not dry:
        backup(path)
        doc.save(path)
        print("talk track saved (%d changes)" % changed)
    return changed


def beat_paragraphs(beat):
    out = [("who", beat["who"]), ("say", beat["say"]), ("label", "SHOW")]
    out += [("bullet", b) for b in beat.get("show", [])]
    if beat.get("point"):
        out.append(("label", "POINT AT"))
        out += [("bullet", b) for b in beat["point"]]
    return out


def find_templates(paras):
    """Who/Say lines with a bold lead-in run plus a normal run; a bullet whose first run is not bold."""
    t = {}
    for p in paras:
        if "who" not in t and p.text.startswith("Who:") and len(p.runs) >= 2:
            t["who"] = p
        elif "say" not in t and p.text.startswith("Say:") and len(p.runs) >= 2:
            t["say"] = p
        elif "label" not in t and p.text.strip() == "SHOW":
            t["label"] = p
        elif ("bullet" not in t and p.style is not None and p.style.name == "List Paragraph" and p.runs
              and not p.runs[0].bold):
            t["bullet"] = p
    return t


def set_lead_text(p_el, text):
    """'Who:  rest' or 'Say: rest': the lead-in goes in the first (bold) run, the rest in the second run."""
    lead, rest = text.split(":", 1)
    pad = rest[:len(rest) - len(rest.lstrip(" "))]
    runs = p_el.findall(W + "r")
    if len(runs) < 2:
        set_element_text(p_el, text)
        return
    for r in runs[2:]:
        p_el.remove(r)
    for run, value in ((runs[0], lead + ":" + pad), (runs[1], rest.lstrip(" "))):
        ts = run.findall(W + "t")
        for t in ts[1:]:
            run.remove(t)
        ts[0].text = value
        ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def add_styled_row(table, texts):
    """Append a copy of the table's first body row that carries the document's cell borders, so new rows match."""
    styled = next(r for r in table.rows[1:] if r._tr.xpath("./w:tc[1]/w:tcPr/w:tcBorders"))
    new = copy.deepcopy(styled._tr)
    table._tbl.append(new)
    row = table.rows[-1]
    for cell, text in zip(row.cells, texts):
        set_cell(cell, text)
    return row


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def set_element_text(p_el, text):
    runs = p_el.findall(W + "r")
    for r in runs[1:]:
        p_el.remove(r)
    if runs:
        ts = runs[0].findall(W + "t")
        for t in ts[1:]:
            runs[0].remove(t)
        if ts:
            ts[0].text = text
            ts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        else:
            import lxml.etree as et
            t = et.SubElement(runs[0], W + "t")
            t.text = text
    else:
        import lxml.etree as et
        r = et.SubElement(p_el, W + "r")
        t = et.SubElement(r, W + "t")
        t.text = text


def set_paragraph(p, text):
    set_element_text(p._element, text)


def set_cell(cell, text):
    ps = cell.paragraphs
    set_paragraph(ps[0], text)
    for q in ps[1:]:
        q._element.getparent().remove(q._element)


def fill_live(cfg):
    sys.path.insert(0, os.path.join(ROOT, "tools"))
    import ztrest
    s = ztrest.Splunk()
    st = s.kv_get("zt_demo_state", "global") or {}
    live = cfg.setdefault("live", {})
    t0 = float(st.get("plan_t0") or 0)
    if t0:
        import time
        live["t0"] = time.strftime("%Y-%m-%d %H:%M:%SZ", time.gmtime(t0))
        rows = s.search('search index=notable source="ZT - Workload Exceeded Risk Threshold on Protected-Path Signals - Rule" earliest=%d | head 1 | table _time event_id' % int(t0), earliest=int(t0), latest="now")
        if rows:
            live["finding_at"] = rows[0]["_time"]
        reqs = s.kv_list("zt_enforcement_requests")
        reqs = [r for r in reqs if float(r.get("requested_epoch") or 0) >= t0]
        if reqs:
            r = reqs[-1]
            live["investigation_id"] = r.get("investigation_id", "")
            for k in ("approved", "applied", "verified"):
                if r.get(k + "_epoch"):
                    live[k + "_at"] = time.strftime("%H:%M:%S", time.gmtime(float(r[k + "_epoch"])))
    with open(os.path.join(HERE, "collateral.yaml"), "w") as fh:
        yaml.safe_dump(cfg, fh, sort_keys=False, allow_unicode=True, width=200)
    print("live section filled: %s" % live)


def main(argv):
    cfg = yaml.safe_load(open(os.path.join(HERE, "collateral.yaml")))
    dry = "--dry-run" in argv
    if "--live" in argv:
        fill_live(cfg)
    n = sync_deck(cfg, dry)
    m = sync_talk_track(cfg, dry)
    print("done: deck %d change(s), talk track %d change(s)%s" % (n, m, " (dry run)" if dry else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
