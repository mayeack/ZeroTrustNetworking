#!/usr/bin/env python3
"""Verify the zero trust collateral against the 24 Sep 2026 review (see the handoff checklist in the OneDrive folder).

Usage: python3 docs/collateral/verify_collateral.py [--stack]
  --stack  also read the Secure Networking Essentials setup state from the Cloud stack (needs local/env).
Prints PASS / FAIL / WARN lines and exits 1 if any FAIL. WARN lines are stale or over-claiming phrases to fix
(sections S and F of the checklist); they do not fail the run.
"""
import os
import re
import subprocess
import sys

import docx
from pptx import Presentation

OD = "/Users/myeack/Library/CloudStorage/OneDrive-Cisco/Projects/Sales Plays/Zero Trust Networking/"
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OI = OD + "Zero Trust Networking - First-Call Deck (One Incident, End to End).pptx"
FC = OD + "Zero Trust Networking - First-Call Deck.pptx"
TT = OD + "Zero Trust Networking - Demo Talk Track.docx"
PDFS = {"leave": OD + "Zero Trust Networking - Customer Leave-Behind.pdf",
        "cosell": OD + "Zero Trust Networking - Internal Co-Sell One-Pager.pdf",
        "oneinc": OD + "Zero Trust Networking - One Incident End to End - One-Pager.pdf"}
ASK = "one cluster, one protected store in audit mode and one Nexus leaf pair"
BANNED = re.compile(r"\b(sample|mock|mockup|fake|synthetic|illustrative)\b", re.I)
ACCOUNT = re.compile(r"SpaceX|\bxAI\b|Starbase|Memphis|Hawthorne|Colossus")
OI_TITLES = {1: "One Incident, End to End", 3: "Three initiatives", 4: "One incident, five steps", 5: "What you'll see live",
             6: "How the pieces connect", 7: "Behind the scenes", 8: "The build farm", 9: "The training checkpoint store",
             10: "Cilium: identity and audit mode", 11: "Hubble and Tetragon", 12: "Enforcement points",
             13: "What ships, and what we build for you", 14: "Live in Splunk", 15: "Step 2 in Splunk", 16: "Step 3: two signals",
             17: "Step 3: the agent", 18: "Step 4", 19: "Step 5", 20: "What each step is worth", 21: "From this demo",
             22: "Thank you"}
STALE = [  # (label, regex) -> WARN, see checklist sections S and F
    ("S1 Runtime Detections now shows SIGKILL alerts (1.0.7)", r"Runtime Detections reads 0|sends no Tetragon policy alerts"),
    ("S2 Tetragon alert event already done (1.0.7)", r"Tetragon alert event"),
    ("S3 ESCU fields are calculated fields in 1.0.7, not a to-do", r"need field aliases|with field aliases|field aliases for the Splunk"),
    ("F1 Hypershield enforces on the Smart Switch; BlueField DPU is Hybrid Mesh Firewall",
     r"Hypershield rule on the DPU|DPU or smart switch|Hypershield (enforcement )?on (NVIDIA )?BlueField|Hypershield on BlueField|Hybrid Mesh Firewall and Hypershield"),
    ("F2 INV SGT is one-way, NX-OS VXLAN EVPN, hand-assigned, Cilium-only",
     r"Security Group Tags? (travel|propagat\w*|carried)|carries Security Group Tags|Security Group Tag propagation"),
    ("F3 Network Explorer with Isovalent is a preview", r"Network Explorer (runs|on Tetragon)(?![^.]{0,40}preview)"),
]
results = []


def rec(status, msg):
    results.append(status)
    print("%-4s %s" % (status, msg))


def check(cond, msg):
    rec("PASS" if cond else "FAIL", msg)


def slide_text(s):
    out = []
    for sh in s.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if sh.shape_type == 6:
            out += [x.text_frame.text for x in sh.shapes if x.has_text_frame]
    return "\n".join(out)


def pdf_text(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout


# a match is acceptable when its sentence carries the qualifier the fact-check asked for
QUALIFIED = {"F2": re.compile(r"Networking for Virtualization|\bINV\b|EVPN", re.I)}


def sentence_around(text, start, end):
    """The sentence holding a match; line breaks inside a sentence (PDF text) do not end it."""
    left = text.rfind(". ", 0, start)
    right = text.find(". ", end)
    return text[left + 1:right if right != -1 else len(text)]


def warn_stale(label, text):
    for name, pat in STALE:
        for m in re.finditer(pat, text, re.I):
            q = QUALIFIED.get(name.split()[0])
            if q and q.search(sentence_around(text, m.start(), m.end())):
                continue
            rec("WARN", "%s: %s ...%s..." % (name, label, text[max(0, m.start() - 50):m.end() + 40].replace("\n", " ")))


# ---------------------------------------------------------------- decks
oi = Presentation(OI)
check(len(oi.slides) == 22, "One Incident deck has 22 slides (found %d)" % len(oi.slides))
for n, t in OI_TITLES.items():
    if n <= len(oi.slides):
        check(t.lower() in slide_text(oi.slides[n - 1]).lower(), "One Incident slide %d contains '%s'" % (n, t))
check(ASK in slide_text(oi.slides[20]), "One Incident slide 21 carries the harmonised ask")
for n in (3, 13, 21, 22):
    check(oi.slides[n - 1].has_notes_slide and len(oi.slides[n - 1].notes_slide.notes_text_frame.text) > 40,
          "One Incident slide %d has speaker notes" % n)
fc = Presentation(FC)
check(len(fc.slides) == 17, "Original deck still has 17 slides (found %d)" % len(fc.slides))
check(ASK in slide_text(fc.slides[15]), "Original deck slide 16 carries the harmonised ask")
for tag, prs in (("One Incident", oi), ("Original", fc)):
    for i, s in enumerate(prs.slides, 1):
        t = slide_text(s)
        check(not BANNED.search(t), "%s slide %d: no banned words" % (tag, i)) if BANNED.search(t) else None
        check(not ACCOUNT.search(t), "%s slide %d: no account names" % (tag, i)) if ACCOUNT.search(t) else None
        if s.has_notes_slide and ACCOUNT.search(s.notes_slide.notes_text_frame.text):
            rec("WARN", "%s slide %d notes name the account (customer-shared decks must not)" % (tag, i))
        warn_stale("%s slide %d" % (tag, i), t)
        if s.has_notes_slide:
            warn_stale("%s slide %d notes" % (tag, i), s.notes_slide.notes_text_frame.text)
rec("PASS", "Decks: banned-word and account-name scan done (FAIL lines above if any)")

# ---------------------------------------------------------------- talk track
d = docx.Document(TT)
paras = d.paragraphs
beats = [p.text for p in paras if p.style is not None and p.style.name.startswith("Heading 2") and p.text.startswith("Beat ")]
check([b.split(":")[0] for b in beats] == ["Beat %d" % i for i in range(1, 11)], "Talk track has Beat 1 to Beat 10 in order")
body = "\n".join(p.text for p in paras)
refs = [int(x) for x in re.findall(r"Slides? (\d+)", body)] + [int(b) for a, b in re.findall(r"Slides (\d+) (?:to|and) (\d+)", body)]
check(refs and max(refs) <= 22, "Talk track slide references stay within 1-22 (max %s)" % (max(refs) if refs else None))
sec4 = False
bold_bullets, bad_lead = 0, 0
for p in paras:
    if p.text.startswith("4. The talk track"):
        sec4 = True
    elif p.text.startswith("5. Discovery questions"):
        sec4 = False
    if not sec4 or not p.runs:
        continue
    if p.style is not None and p.style.name == "List Paragraph" and p.runs[0].bold:
        bold_bullets += 1
    if p.text.startswith(("Who:", "Say:")) and (len(p.runs) < 2 or p.runs[1].bold):
        bad_lead += 1
check(bold_bullets == 0, "Talk track section 4: no bold bullets (found %d)" % bold_bullets)
check(bad_lead == 0, "Talk track section 4: Who/Say lines have a bold lead-in and a normal remainder (bad %d)" % bad_lead)
unstyled = [(ti, r.cells[0].text[:30]) for ti, t in enumerate(d.tables) for r in t.rows[1:]
            if not r._tr.xpath("./w:tc[1]/w:tcPr/w:tcBorders")]
check(not unstyled, "Talk track tables: every body row carries the table style (unstyled: %s)" % unstyled)
t3 = [r.cells[0].text.strip() for r in d.tables[0].rows]
check("Secure Networking Essentials" in t3, "Talk track section 3 has the Secure Networking Essentials row")
t6 = [r.cells[0].text.strip() for r in d.tables[1].rows]
check("We saw Cisco Secure Workload; why not that?" in t6 and "Do we have to build all these dashboards and detections?" in t6,
      "Talk track section 6 has the two new objections")
t7 = [r.cells[0].text.strip() for r in d.tables[2].rows]
for claim in ("Splunk Security Content: Cisco Isovalent detections", "Splunk Secure Networking Essentials (Beta)",
              "Cisco Enterprise Networking app and add-on", "Cisco DC Networking app",
              "Enterprise Security 8.7 tools in the Splunk MCP Server", "dCloud Cisco Secure Networking demo"):
    check(claim in t7, "Talk track section 7 row: %s" % claim)
check(ASK in body, "Talk track carries the harmonised ask")
check("job class" not in body.lower(), "Talk track no longer puts the job class in the ask")
warn_stale("Talk track body", body)
for ti, t in enumerate(d.tables):
    for r in t.rows:
        warn_stale("Talk track table %d '%s'" % (ti, r.cells[0].text[:30]), " | ".join(c.text for c in r.cells))

# ---------------------------------------------------------------- PDFs
for key, path in PDFS.items():
    info = subprocess.run(["pdfinfo", path], capture_output=True, text=True).stdout
    pages = re.search(r"Pages:\s+(\d+)", info)
    check(pages and pages.group(1) == "1", "%s PDF is one page" % key)
    size = "612 x 792" if key == "cosell" else "792 x 612"
    check(size in info, "%s PDF page size %s" % (key, size))
    t = pdf_text(path)
    if key != "cosell":
        check(not re.search(r"\bbeta\b", t, re.I), "%s PDF (customer) names no beta product" % key)
        check(not BANNED.search(t), "%s PDF (customer) has no banned words" % key)
    check(not ACCOUNT.search(t), "%s PDF names no account" % key)
    warn_stale("%s PDF" % key, t)
check(ASK in re.sub(r"\s+", " ", pdf_text(PDFS["leave"])), "Leave-behind carries the harmonised ask")

# ---------------------------------------------------------------- sync and repo docs
out = subprocess.run([sys.executable, os.path.join(ROOT, "docs/collateral/sync_collateral.py"), "--dry-run"],
                     capture_output=True, text=True, cwd=ROOT).stdout
check("deck 0 change(s), talk track 0 change(s)" in out, "make collateral --dry-run reports 0 changes")
ds = open(os.path.join(ROOT, "docs/demo_script.md")).read()
check("Before the fire: what ships" in ds, "demo_script.md has the 'Before the fire: what ships' section")
check(os.path.getmtime(os.path.join(ROOT, "docs/demo_script.html")) >= os.path.getmtime(os.path.join(ROOT, "docs/demo_script.md")),
      "demo_script.html is at least as new as demo_script.md")
warn_stale("demo_script.md", ds)
warn_stale("collateral.yaml", open(os.path.join(ROOT, "docs/collateral/collateral.yaml")).read())
ART = OD + "Claude Artifacts/"  # review, handoff checklist and TA generating prompt live here since 24 Sep 2026
rv = (ART if os.path.exists(ART + "Zero Trust Networking - Demo Review (dCloud and Splunk Apps).md") else OD) + "Zero Trust Networking - Demo Review (dCloud and Splunk Apps)"
check(os.path.exists(rv + ".md") and os.path.exists(rv + ".html"), "Review .md and .html both exist")
check(os.path.getmtime(rv + ".html") >= os.path.getmtime(rv + ".md"), "Review .html is at least as new as the .md")
warn_stale("Review .md", open(rv + ".md").read())
bk = os.path.join(ROOT, "docs/collateral/backup/2026-09-24")
check(os.path.isdir(bk) and len(os.listdir(bk)) >= 7, "Backups of the 24 Sep originals present (%s)" % bk)

# ---------------------------------------------------------------- stack (optional)
if "--stack" in sys.argv:
    sys.path.insert(0, os.path.join(ROOT, "tools"))
    import base64
    import json
    import ssl
    import urllib.request
    env = {}
    for line in open(os.path.join(ROOT, "local/env")):
        line = line.strip()
        if "=" in line and not line.startswith("#"):
            k, v = line.split("=", 1)
            env[k] = v.strip().strip('"').strip("'")
    auth = "Basic " + base64.b64encode(("%s:%s" % (env["SPLUNK_USER"], env["SPLUNK_PASS"])).encode()).decode()

    def get(path):
        req = urllib.request.Request(env["SPLUNK_URL"].rstrip("/") + path + ("&" if "?" in path else "?") + "output_mode=json",
                                     headers={"Authorization": auth})
        return json.loads(urllib.request.urlopen(req, context=ssl.create_default_context(), timeout=60).read())
    app = "splunk-cisco-dashboards"
    c = get("/servicesNS/nobody/%s/configs/conf-setup/install" % app)["entry"][0]["content"]
    check(c.get("is_configured") in ("1", 1, True), "Secure Networking Essentials setup is_configured=1")
    check(get("/servicesNS/nobody/%s/admin/macros/index_security" % app)["entry"][0]["content"]["definition"] == "index IN (zero_trust)",
          "index_security = index IN (zero_trust)")
    for m in ("index_cloud_native", "index_cybervision", "index_windows"):
        check(get("/servicesNS/nobody/%s/admin/macros/%s" % (app, m))["entry"][0]["content"]["definition"] == "index IN (*)",
              "%s unchanged (index IN (*))" % m)
    check(get("/servicesNS/nobody/cisco-catalyst-app/admin/macros/cisco_catalyst_app_index")["entry"][0]["content"]["definition"] == "index IN (*)",
          "cisco_catalyst_app_index unchanged (index IN (*))")
    cim = get("/servicesNS/nobody/Splunk_SA_CIM/admin/macros?count=0&search=name%3Dcim_*_indexes")["entry"]
    check(all(e["content"]["definition"] == "()" for e in cim), "CIM index allowlists untouched (all '()')")
    check(get("/services/apps/local/searchbase")["entry"][0]["content"].get("disabled") in (False, "0", 0),
          "Searchbase app installed and enabled")

fails = results.count("FAIL")
print("\nSUMMARY: %d PASS, %d FAIL, %d WARN" % (results.count("PASS"), fails, results.count("WARN")))
sys.exit(1 if fails else 0)
